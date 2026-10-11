#!/usr/bin/env python3
"""Export Efi Baron customers: first name, last name, mobile, email.

Only customers with first name + last name + a valid Israeli mobile.
If Phone1 and Phone2 are both valid mobiles, Phone2 is used (newer).
"""

from __future__ import annotations

import importlib.util
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

SRC_DIR = Path("/Users/maorganon/Downloads/EfiBaronOld-sys")
CUST_PATH = SRC_DIR / "CustList.xlsx"
OUT_PATH = SRC_DIR / "EfiBaron-contacts.xlsx"

HELPERS = Path(__file__).with_name("build-efi-baron-formula-notes.py")
spec = importlib.util.spec_from_file_location("efi_notes", HELPERS)
efi = importlib.util.module_from_spec(spec)
spec.loader.exec_module(efi)

HEADER_FILL = PatternFill("solid", fgColor="1F4E5F")
HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=11)
CELL_FONT = Font(name="Arial", size=11)
BAND_A = PatternFill("solid", fgColor="FFFFFF")
BAND_B = PatternFill("solid", fgColor="E8EEF4")
THIN = Border(
    left=Side(style="thin", color="D0D7DE"),
    right=Side(style="thin", color="D0D7DE"),
    top=Side(style="thin", color="D0D7DE"),
    bottom=Side(style="thin", color="D0D7DE"),
)
RIGHT = Alignment(horizontal="right", vertical="center", readingOrder=2)
LEFT = Alignment(horizontal="left", vertical="center")


def pick_mobile(phone1: str, phone2: str) -> str | None:
    k1, k2 = efi.phone_kind(phone1), efi.phone_kind(phone2)
    n1, n2 = efi.national_phone(phone1), efi.national_phone(phone2)
    m1 = n1 if k1 == "mobile" else None
    m2 = n2 if k2 == "mobile" else None
    if m1 and m2:
        return m2
    return m2 or m1


def clean_email(raw) -> str:
    s = efi.clean(raw)
    if not s or "@" not in s:
        return ""
    return s.strip()


def load_rows() -> list[list[str]]:
    wb = load_workbook(CUST_PATH, data_only=True)
    ws = wb.active
    rows = []
    for r in range(2, ws.max_row + 1):
        first = efi.clean(ws.cell(r, 2).value)
        last = efi.clean(ws.cell(r, 3).value)
        if not first or not last:
            continue
        mobile = pick_mobile(efi.clean(ws.cell(r, 7).value), efi.clean(ws.cell(r, 8).value))
        if not mobile:
            continue
        rows.append([first, last, mobile, clean_email(ws.cell(r, 11).value)])
    wb.close()
    rows.sort(key=lambda x: (x[1], x[0], x[2]))
    return rows


def main() -> None:
    rows = load_rows()
    wb = Workbook()
    ws = wb.active
    ws.title = "לקוחות"
    ws.sheet_view.rightToLeft = True
    ws.freeze_panes = "A2"
    headers = ["שם פרטי", "שם משפחה", "מספר טלפון", "אימייל"]
    widths = [18, 22, 16, 36]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(1, col, header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN
        ws.column_dimensions[get_column_letter(col)].width = widths[col - 1]
    ws.row_dimensions[1].height = 26
    with_email = 0
    for i, row in enumerate(rows, 2):
        fill = BAND_A if i % 2 == 0 else BAND_B
        if row[3]:
            with_email += 1
        for c, value in enumerate(row, 1):
            cell = ws.cell(i, c, value or None)
            cell.font = CELL_FONT
            cell.border = THIN
            cell.fill = fill
            cell.alignment = LEFT if c == 4 else RIGHT
    ws.auto_filter.ref = f"A1:D{max(1, len(rows) + 1)}"
    wb.save(OUT_PATH)
    print(f"wrote {OUT_PATH}")
    print(f"  לקוחות: {len(rows)}")
    print(f"  עם אימייל: {with_email}")


if __name__ == "__main__":
    main()
