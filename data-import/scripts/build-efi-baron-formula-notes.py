#!/usr/bin/env python3
"""Convert Efi Baron legacy CustList + CustColor into formula notes.

Output workbook:
  - רשמי
  - אימות לקוחות
  - בדיקה
  - סיכום
"""

from __future__ import annotations

import re
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

SRC_DIR = Path("/Users/maorganon/Downloads/EfiBaronOld-sys")
CUST_PATH = SRC_DIR / "CustList.xlsx"
COLOR_PATH = SRC_DIR / "CustColor.xlsx"
OUT_PATH = SRC_DIR / "EfiBaron-formula-notes.xlsx"

MOBILE_PREFIXES = ("050", "052", "053", "054", "055", "057", "058")
OXY_PERCENTS = {
    "1.2", "1.5", "1.8", "2.4", "2.7", "2.8", "3", "3.75", "4", "4.3", "4.5",
    "4.6", "6", "6.6", "7.5", "9", "9.8", "12",
}
OXY_VOLUMES = {"10", "20", "30", "40"}

FRACTIONS = {
    "1/8": "שמינית טיובה",
    "1\\8": "שמינית טיובה",
    "1/4": "רבע טיובה",
    "1\\4": "רבע טיובה",
    "1/2": "חצי טיובה",
    "1\\2": "חצי טיובה",
    "3/4": "שלושת רבעי טיובה",
    "3\\4": "שלושת רבעי טיובה",
}

HEADER_WORDS = {
    "שטיפה": "שטיפה",
    "שטיפות": "שטיפה",
    "גוונים": "גוונים",
    "גבות": "גבות",
    "פיגמנטציה": "פיגמנטציה",
    "אורכים": "אורכים",
    "אורך": "אורכים",
    "שמן": "שמן",
    "סיבוב שני": "סיבוב שני",
    "סיבוב 2": "סיבוב שני",
}

SERIES_WORDS = {
    "inoa": "INOA",
    "אינואה": "INOA",
    "איואה": "INOA",
    "אינאה": "INOA",
    "וולה": "Wella",
    "wella": "Wella",
    "דורה": "Dore",
    "indola": "Indola",
    "אינדו": "Indola",
    "אינדולה": "Indola",
}

CLIENT_NOTE_RE = re.compile(r"^(רגישות\+?|חומר רגישו.*|\*{2,})$")
HADASH_RE = re.compile(r"חדש")
YASHAN_RE = re.compile(r"^ישן$")
PRICE_RE = re.compile(
    r"(תספורת|קרטין|\bפן\b|פריט|\d+\s*%|ש\"?ח|צ\s*\+\s*ת|צבע\s*\+?\s*(חצי|גוו)?)",
    re.I,
)
INSTRUCTION_RE = re.compile(
    r"(דק|דקות|להשאיר|שורש|קצוות|לא לגעת|להוסיף|לצבע הקיים|רק ב|בלי פן|אחרי צבע|נקודתי)",
)
PRODUCT_RE = re.compile(
    r"(שמן|בוסטר|קליר|קלייר|שטין|שטיין|פרונ|פונס|מרון|ויולין|רוז׳?|rouge|cuivre|dore|chatain|chatin|blond|בלונד|plex|olaplex)",
    re.I,
)
SHADE_CORE_RE = re.compile(
    r"^([cCrR])?\s*(\d{1,2}(?:[./,]\d{1,3})?)(?:\s*([cCrR]|KD|kd|SU|Su|S|s))?(.*)$"
)
COMPACT_SHADE_RE = re.compile(r"^(\d{3,4})([A-Za-z+]*)$")

SERVICE_PRICE_LABELS = (
    ("תספורת", "תספורת"),
    ("תספרת", "תספורת"),
    ("קרטין", "קרטין"),
    ("החלקה", "החלקה"),
    ("שטיפות", "שטיפה"),
    ("שטיפה", "שטיפה"),
    ("גוונים", "גוונים"),
    ("צבע", "צבע"),
    ("פן", "פן"),
)


def is_ils_price_number(token: str) -> bool:
    if not re.fullmatch(r"\d{2,4}", str(token or "").strip()):
        return False
    return 40 <= int(token) <= 2500


def parse_service_price(shade: str, amt: str) -> tuple[str, int] | None:
    """'גוונים 300' is 300 ILS for highlights, not shade 3.00."""
    s, a = clean(shade), clean(amt)
    formula_qty = a in FRACTIONS or a in {"שלם", "מלא", "מעט", "1/8", "1/4", "1/2", "3/4"}
    blob = s.replace(":", " ")
    for he, label in SERVICE_PRICE_LABELS:
        m = re.fullmatch(rf"{re.escape(he)}\s*(\d{{2,4}})", blob)
        if m and is_ils_price_number(m.group(1)):
            return label, int(m.group(1))
        m = re.fullmatch(rf"(\d{{2,4}})\s*{re.escape(he)}", blob)
        if m and is_ils_price_number(m.group(1)):
            return label, int(m.group(1))
        if s == he and not formula_qty and is_ils_price_number(a):
            return label, int(a)
    return None

HEADER_FILL = PatternFill("solid", fgColor="1F4E5F")
HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=11)
CELL_FONT = Font(name="Arial", size=11)
WARN_FILL = PatternFill("solid", fgColor="FFF4D5")
CUST_FILL_A = PatternFill("solid", fgColor="FFFFFF")
CUST_FILL_B = PatternFill("solid", fgColor="E8EEF4")
THIN = Border(
    left=Side(style="thin", color="D0D7DE"),
    right=Side(style="thin", color="D0D7DE"),
    top=Side(style="thin", color="D0D7DE"),
    bottom=Side(style="thin", color="D0D7DE"),
)
WRAP = Alignment(horizontal="right", vertical="top", wrap_text=True, readingOrder=2)
RIGHT = Alignment(horizontal="right", vertical="center", readingOrder=2)
CENTER = Alignment(horizontal="center", vertical="center")
RLM = "\u200F"
LINE_HEIGHT = 16
ROW_HEIGHT_CAP = 480


def visual_line_count(text: str, col_width: int) -> int:
    if not text:
        return 1
    chars_per_line = max(10, int(col_width * 1.05))
    total = 0
    for para in str(text).split("\n"):
        length = max(1, len(para))
        total += max(1, (length + chars_per_line - 1) // chars_per_line)
    return total


def fitted_row_height(row: list, wrap_cols: set[int], widths: list[int]) -> float:
    lines = 1
    for c_i in wrap_cols:
        value = row[c_i - 1] if c_i - 1 < len(row) else None
        if not value:
            continue
        width = widths[c_i - 1] if c_i - 1 < len(widths) else 40
        lines = max(lines, visual_line_count(str(value), width))
    return min(ROW_HEIGHT_CAP, max(22, lines * LINE_HEIGHT + 6))


def clean(value) -> str:
    if value is None:
        return ""
    text = str(value).replace("\u00a0", " ").strip()
    text = " ".join(text.split())
    if text in {"-", '""', "None"}:
        return ""
    return text


def digits_only(value: str) -> str:
    return re.sub(r"\D", "", value or "")


def national_phone(raw: str) -> str:
    d = digits_only(raw)
    if d.startswith("972"):
        d = d[3:]
    if d and not d.startswith("0"):
        d = "0" + d
    return d


def is_valid_mobile(national: str) -> bool:
    return len(national) == 10 and national.startswith(MOBILE_PREFIXES)


def phone_kind(raw: str) -> str:
    d = national_phone(raw)
    if not d:
        return "empty"
    if is_valid_mobile(d):
        return "mobile"
    if d.startswith("051") or d.startswith("056") or d.startswith("059"):
        return "mobile_uncertain"
    if d.startswith(("02", "03", "04", "08", "09", "072", "073", "074", "076", "077", "079")):
        return "landline"
    if d.startswith("05"):
        return "mobile_uncertain"
    return "other"


def pick_mobile(phone1: str, phone2: str) -> tuple[str | None, list[str]]:
    flags = []
    k1, k2 = phone_kind(phone1), phone_kind(phone2)
    n1, n2 = national_phone(phone1), national_phone(phone2)
    mobiles = []
    if k1 == "mobile":
        mobiles.append(("1", n1))
    if k2 == "mobile":
        mobiles.append(("2", n2))
    if len(mobiles) >= 2:
        flags.append("שני ניידים תקינים — נלקח Phone1")
        return mobiles[0][1], flags
    if mobiles:
        return mobiles[0][1], flags
    return None, flags


def format_amount(raw: str) -> tuple[str | None, list[str], list[str]]:
    """Return (display, flags, corrections)."""
    flags, corr = [], []
    s = clean(raw)
    if not s:
        return None, ["כמות חסרה"], corr

    compact = s.replace(" ", "")
    plus = compact.endswith("+")
    minus = compact.endswith("-")
    core = compact[:-1] if plus or minus else compact
    suffix = "+" if plus else ("-" if minus else "")

    if core in FRACTIONS:
        out = FRACTIONS[core] + suffix
        return out, flags, corr
    if s in {"שלם", "מלא", "בול"} or core == "שלם":
        out = "טיובה שלמה" + ("+" if plus or s == "שלם+" else "")
        if s not in {"טיובה שלמה"}:
            corr.append(f"{s} → {out}")
        return out, flags, corr
    if core == "1":
        flags.append("כמות 1 פורשה כטיובה שלמה")
        corr.append(f"{s} → טיובה שלמה")
        return "טיובה שלמה", flags, corr

    gram = re.fullmatch(r"(\d+(?:\.\d+)?)\s*(גר'|גרם|גר|ג|gr|g|G)", s, re.I)
    if gram:
        out = f"{gram.group(1)} גרם"
        if out != s:
            corr.append(f"{s} → {out}")
        return out, flags, corr

    cap = re.fullmatch(r"(?:פקק\s*)?(\d+)\s*פקק", s)
    if cap or re.fullmatch(r"\d+פקק", compact):
        n = cap.group(1) if cap else re.match(r"(\d+)", compact).group(1)
        out = f"{n} פקק"
        if out != s:
            corr.append(f"{s} → {out}")
        return out, flags, corr
    if s == "פקק":
        return "פקק", flags, corr

    if re.fullmatch(r"\d+(?:\.\d+)?", s):
        flags.append("כמות מספרית בלי יחידה")
        return s, flags, corr

    keep = {"מעט", "מעט-", "מעט -", "קצת", "טיפה", "נגיעה", "כפית", "שורה", "מלא"}
    if s in keep or s.startswith("מעט"):
        return s, flags, corr

    flags.append("כמות לא סטנדרטית")
    return s, flags, corr


def is_loreal_high_lift_name(raw: str) -> bool:
    """L'Oréal high-lift / Super Lightener names like 900S, 901s, 900 — not 9.01."""
    t = clean(raw).replace(" ", "").replace("/", "").replace("\\", "")
    return bool(re.fullmatch(r"9(00|01|02|03)s?", t, re.I))


def glue_shade_letters(s: str) -> str:
    """Keep C/R/KD attached to the number so RTL Excel does not split shade from amount.

    '7.1 c' / 'c 7.1' / '8 R' → '7.1C' / '8R'
    """
    t = re.sub(r"\s+", " ", clean(s))
    m = re.fullmatch(
        r"([cCrR])\s*(\d{1,2}(?:[./]\d{1,3})?)\s*([cCrR]|KD|kd)?",
        t,
    )
    if m:
        num = m.group(2)
        letters = "".join(x.upper() for x in (m.group(1), m.group(3) or "") if x)
        if len(letters) >= 2 and letters[0] == letters[-1]:
            letters = letters[0]
        return f"{num}{letters}"
    m = re.fullmatch(
        r"(\d{1,2}(?:[./]\d{1,3})?)\s*([cCrR]|KD|kd)",
        t,
    )
    if m:
        return f"{m.group(1)}{m.group(2).upper()}"
    m = re.fullmatch(r"(\d{1,2}(?:[./]\d{1,3})?)([cCrR]|KD|kd)", t, re.I)
    if m:
        return f"{m.group(1)}{m.group(2).upper()}"
    return t


def pair_line(left: str, right: str | None) -> str:
    """Visual RTL: shade on the right, space-dash-space, then amount."""
    if not left:
        return right or ""
    line = f"{left} - {right}" if right else left
    # Only pin digit/latin shades to the right. Hebrew-first lines (חמצן) already sit correctly;
    # a leading RLM there swaps the % with the ml amount.
    if re.match(r"[0-9A-Za-z]", left) and not line.startswith(RLM):
        return f"{RLM}{line}"
    return line


def normalize_shade(raw: str) -> tuple[str, list[str], list[str]]:
    flags, corr = [], []
    s = clean(raw)
    if not s:
        return "", flags, corr
    original = s

    s = s.replace(",", ".")
    if is_loreal_high_lift_name(s):
        return original, flags, corr

    compact = COMPACT_SHADE_RE.fullmatch(s.replace(" ", ""))
    if compact:
        num, tail = compact.group(1), compact.group(2)
        if num == "000":
            s = f"000{tail}" if tail else "000"
        elif len(num) == 3:
            s = f"{num[0]}.{num[1:]}{tail}"
            flags.append("גוון קומפקטי הומר")
            corr.append(f"{original} → {s}")
        elif len(num) == 4 and num.startswith(("10", "11", "12")):
            s = f"{num[:2]}.{num[2:]}{tail}"
            flags.append("גוון קומפקטי הומר")
            corr.append(f"{original} → {s}")

    glued = glue_shade_letters(s)
    if glued != s:
        if corr:
            corr[-1] = f"{original} → {glued}"
        s = glued
    elif s != original and "/" not in original and not corr:
        corr.append(f"{original} → {s}")
    return s, flags, corr


def looks_like_shade(text: str) -> bool:
    s = clean(text)
    if not s:
        return False
    if CLIENT_NOTE_RE.match(s) or HADASH_RE.search(s) or YASHAN_RE.match(s):
        return False
    if s.lower() in SERIES_WORDS or clean(s) in HEADER_WORDS:
        return False
    if PRICE_RE.search(s) and not re.search(r"\d{1,2}[./]\d", s):
        return False
    core = s.replace(" ", "")
    if re.fullmatch(r"[cCrR]?\d{1,2}([./,]\d{1,3})?[A-Za-z+]*", core):
        return True
    if re.fullmatch(r"\d{3,4}[A-Za-z]*", core):
        return True
    if SHADE_CORE_RE.match(s):
        return bool(re.search(r"\d", s))
    return False


def split_embedded_header(shade: str) -> tuple[str, str | None]:
    s = clean(shade)
    m = re.match(r"^(פיגמנטציה)\s*(\d+[./]?\d*)?$", s)
    if m:
        rest = m.group(2) or ""
        return rest, "פיגמנטציה"
    for key, label in (
        ("שטיפה", "שטיפה"),
        ("גבות", "גבות"),
        ("אורכים", "אורכים"),
        ("גוונים", "גוונים"),
    ):
        if key in s and s != key:
            left = s.replace(key, " ").strip()
            if is_ils_price_number(left.replace(" ", "")):
                return s, None
            if looks_like_shade(left) or re.search(r"\d", left):
                return left, label
        if s.endswith(" ש") and looks_like_shade(s[:-2]):
            return s[:-2].strip(), "שטיפה"
    if re.search(r"\sגב$", s) or s.endswith(" גב"):
        left = re.sub(r"\sגב$", "", s).strip()
        if looks_like_shade(left) or re.search(r"\d", left):
            return left, "גבות"
    if s.endswith(" ש") and re.search(r"\d", s[:-2]):
        return s[:-2].strip(), "שטיפה"
    return s, None


def classify_text_row(shade: str, amt: str) -> str | None:
    blob = f"{shade} {amt}".strip()
    if not blob:
        return "empty"
    if HADASH_RE.search(blob) and not looks_like_shade(shade):
        return "hadash"
    if YASHAN_RE.match(shade):
        return "yashan"
    if CLIENT_NOTE_RE.match(shade) or CLIENT_NOTE_RE.match(amt):
        return "client_note"
    if shade in HEADER_WORDS and not looks_like_shade(shade):
        if shade == "שמן" and amt:
            return "product"
        return "header"
    if re.search(r"שטיפה", shade) and not re.search(r"\d", shade):
        return "header"
    if shade.lower() in SERIES_WORDS or clean(shade).lower() in SERIES_WORDS:
        return "series"
    if INSTRUCTION_RE.search(blob) and not looks_like_shade(shade):
        return "instruction"
    if PRICE_RE.search(blob) and not looks_like_shade(shade):
        return "price"
    if PRODUCT_RE.search(shade) and not looks_like_shade(re.sub(r"(שמן|בוסטר).*$", "", shade)):
        # product line like "שמן 20 גר'" or "שטין קליר"
        if looks_like_shade(split_embedded_header(shade)[0]) and "שמן" not in shade:
            return None
        return "product"
    return None


def format_oxy(percent: str, amount: str) -> tuple[str | None, list[str], list[str]]:
    flags, corr = [], []
    p, a = clean(percent), clean(amount)
    if not p and not a:
        return None, flags, corr
    if not p:
        flags.append("חמצן בלי ריכוז")
        label = f'חמצן{RLM} - {a} מ"ל' if a else "חמצן"
        return label, flags, corr
    p_num = p.replace("%", "").replace("vol", "").replace("VOL", "").strip()
    if p_num in OXY_PERCENTS:
        core = f"חמצן {p_num}%"
    elif p_num in OXY_VOLUMES:
        core = f"חמצן {p_num} vol"
        flags.append("חמצן נראה כ-volume לא כאחוז")
        corr.append(f"{p} → {p_num} vol")
    else:
        core = f"חמצן {p}"
        flags.append("ריכוז חמצן חריג")
    if a:
        # RLM after % keeps 70 מ"ל from merging into 9%
        return f'{core}{RLM} - {a} מ"ל', flags, corr
    flags.append("חמצן בלי כמות")
    return core, flags, corr


def raw_line(row: dict) -> str:
    bits = []
    if row.get("f") not in (None, ""):
        bits.append(f"F{row['f']}")
    for key, label in (("shade", ""), ("amt", "כמות"), ("oxy", "חמצן"), ("oxyAmt", "מ״ל")):
        val = clean(row.get(key))
        if val:
            bits.append(f"{label} {val}".strip() if label else val)
    return " | ".join(bits)


def classify_row(row: dict) -> dict:
    shade_raw, amt_raw = clean(row["shade"]), clean(row["amt"])
    oxy_raw, oxy_amt_raw = clean(row["oxy"]), clean(row["oxyAmt"])
    empty = not any([shade_raw, amt_raw, oxy_raw, oxy_amt_raw])
    item = {
        "kind": "empty" if empty else "unknown",
        "shade": "",
        "shade_raw": shade_raw,
        "amount": None,
        "amount_raw": amt_raw,
        "oxy_line": None,
        "header": None,
        "series": None,
        "text": "",
        "flags": [],
        "corr": [],
        "raw": row,
    }
    if empty:
        return item

    priced = parse_service_price(shade_raw, amt_raw)
    if priced:
        label, shekels = priced
        item["kind"] = "price"
        item["text"] = f'{label} - {shekels} ש"ח'
        item["flags"].append('מחיר שירות בש"ח ולא גוון')
        item["corr"].append(f'{shade_raw} {amt_raw} → {label} - {shekels} ש"ח'.strip())
        return item

    kind = classify_text_row(shade_raw, amt_raw)
    shade_left, embedded_header = split_embedded_header(shade_raw)
    if embedded_header and looks_like_shade(shade_left):
        kind = None
        item["header"] = embedded_header
        item["flags"].append("כותרת פוצלה מתוך שורת גוון")
        item["corr"].append(f"{shade_raw} → גוון {shade_left} + כותרת {embedded_header}")
        shade_raw = shade_left

    if kind == "hadash":
        item["kind"] = "hadash"
        return item
    if kind == "yashan":
        item["kind"] = "yashan"
        item["text"] = "ישן"
        return item
    if kind == "client_note":
        item["kind"] = "client_note"
        if "רגישות+" in shade_raw:
            item["text"] = "**רגישות+**"
        elif "רגישות" in shade_raw or "רגישות" in amt_raw or "רגישו" in shade_raw:
            item["text"] = "**רגישות**"
        else:
            item["text"] = "***"
        return item
    if kind == "header":
        item["kind"] = "header"
        item["header"] = HEADER_WORDS.get(shade_raw, shade_raw)
        return item
    if kind == "series":
        item["kind"] = "series"
        item["series"] = SERIES_WORDS.get(shade_raw.lower(), shade_raw)
        if amt_raw:
            amt, af, ac = format_amount(amt_raw)
            item["amount"] = amt
            item["flags"].extend(af)
            item["corr"].extend(ac)
            item["kind"] = "product"
            item["shade"] = item["series"]
        return item
    if kind == "instruction":
        item["kind"] = "instruction"
        item["text"] = " ".join(x for x in (shade_raw, amt_raw) if x)
        item["flags"].append("הנחיית עבודה לפי הקשר")
        return item
    if kind == "price":
        item["kind"] = "price"
        item["text"] = " ".join(x for x in (shade_raw, amt_raw) if x)
        return item
    if kind == "product":
        item["kind"] = "product"
        name = shade_raw
        amt_source = amt_raw
        m = re.match(r"^(שמן)\s+(.+)$", shade_raw)
        if m:
            name = m.group(1)
            if not amt_source:
                amt_source = m.group(2)
        amt, af, ac = format_amount(amt_source) if amt_source else (None, [], [])
        if amt_source and not amt:
            amt, af, ac = amt_source, [], []
        item["shade"] = name
        item["amount"] = amt
        item["flags"].extend(af)
        item["corr"].extend(ac)
        oxy, of, oc = format_oxy(oxy_raw, oxy_amt_raw)
        item["oxy_line"] = oxy
        item["flags"].extend(of)
        item["corr"].extend(oc)
        return item

    if looks_like_shade(shade_raw):
        item["kind"] = "shade"
        shade, sf, sc = normalize_shade(shade_raw)
        item["shade"] = shade
        item["flags"].extend(sf)
        item["corr"].extend(sc)
        if amt_raw:
            amt, af, ac = format_amount(amt_raw)
            item["amount"] = amt
            item["flags"].extend(af)
            item["corr"].extend(ac)
        else:
            item["flags"].append("גוון בלי כמות")
        oxy, of, oc = format_oxy(oxy_raw, oxy_amt_raw)
        item["oxy_line"] = oxy
        item["flags"].extend(of)
        item["corr"].extend(oc)
        return item

    if oxy_raw or oxy_amt_raw:
        item["kind"] = "oxidant"
        oxy, of, oc = format_oxy(oxy_raw, oxy_amt_raw)
        item["oxy_line"] = oxy
        item["flags"].extend(of)
        item["corr"].extend(oc)
        if shade_raw or amt_raw:
            item["flags"].append("שורה עם חמצן וטקסט לא מזוהה")
            item["text"] = " ".join(x for x in (shade_raw, amt_raw) if x)
        return item

    if shade_raw or amt_raw:
        item["kind"] = "unknown"
        item["text"] = " ".join(x for x in (shade_raw, amt_raw) if x)
        item["flags"].append("טקסט לא מזוהה — שויך לפי הקשר")
        oxy, of, oc = format_oxy(oxy_raw, oxy_amt_raw)
        item["oxy_line"] = oxy
        item["flags"].extend(of)
        item["corr"].extend(oc)
        return item
    return item


def attach_series(items: list[dict]) -> None:
    series_at = [i for i, it in enumerate(items) if it["kind"] == "series"]
    current = None
    for it in items:
        if it["kind"] == "series":
            current = it["series"]
            continue
        if it["kind"] == "header":
            continue
        if it["kind"] == "shade" and current and not it.get("series"):
            it["series"] = current
    if series_at:
        for idx in series_at:
            series = items[idx]["series"]
            # backward fill until header/series/shade-with-series
            j = idx - 1
            while j >= 0:
                it = items[j]
                if it["kind"] in {"header", "series"}:
                    break
                if it["kind"] == "shade" and not it.get("series"):
                    it["series"] = series
                    it["flags"].append("סדרה שויכה לפי קרבה")
                j -= 1


def infer_missing_amounts(items: list[dict]) -> None:
    shade_items = [it for it in items if it["kind"] == "shade"]
    known = [it["amount"] for it in shade_items if it["amount"] and "כמות חסרה" not in it["flags"] and "גוון בלי כמות" not in it["flags"]]
    if not known:
        return
    # only infer if neighbors share one amount
    for i, it in enumerate(shade_items):
        if it["amount"] or "גוון בלי כמות" not in it["flags"]:
            continue
        neighbors = []
        if i:
            neighbors.append(shade_items[i - 1]["amount"])
        if i + 1 < len(shade_items):
            neighbors.append(shade_items[i + 1]["amount"])
        neighbors = [n for n in neighbors if n]
        if len(set(neighbors)) == 1:
            it["amount"] = neighbors[0]
            it["flags"].append("כמות הוסקה מהשורות הסמוכות")
            it["corr"].append(f"(חסר) → {neighbors[0]}")
            it["flags"] = [f for f in it["flags"] if f != "גוון בלי כמות"]


def formula_is_price_like(items: list[dict]) -> bool:
    kinds = [it["kind"] for it in items if it["kind"] not in {"empty", "hadash", "yashan", "client_note"}]
    if not kinds:
        return False
    if any(k == "shade" for k in kinds):
        return False
    return any(k == "price" for k in kinds) or all(k in {"price", "unknown", "instruction"} for k in kinds)


def formula_is_color(items: list[dict]) -> bool:
    return any(it["kind"] == "shade" for it in items)


def render_formula(items: list[dict], client_notes: list[str], price_like: bool) -> tuple[str, list[str], list[str]]:
    flags, corr = [], []
    for it in items:
        flags.extend(it["flags"])
        corr.extend(it["corr"])

    lines: list[str] = []
    if client_notes:
        lines.extend(client_notes)
        lines.append("")
    if price_like:
        lines.append("לבדיקה - כנראה מחיר/שירות ולא פורמולת צבע")
        lines.append("")
        flags.append("כנראה מחיר/שירות ולא פורמולת צבע")

    sections = [{"header": None, "shades": [], "products": [], "oxy": []}]
    instructions: list[str] = []
    unknowns: list[str] = []

    def cur():
        return sections[-1]

    def start_section(header):
        if not (cur()["header"] or cur()["shades"] or cur()["products"] or cur()["oxy"]):
            cur()["header"] = header
            return
        sections.append({"header": header, "shades": [], "products": [], "oxy": []})

    for it in items:
        kind = it["kind"]
        if kind in {"empty", "hadash", "yashan", "client_note", "series"}:
            continue
        if kind == "header":
            start_section(it["header"] or it.get("text") or it.get("shade_raw") or "שטיפה")
            continue
        if kind == "instruction":
            instructions.append(it["text"])
            continue
        if kind == "price":
            cur()["products"].append(it["text"])
            continue
        if kind == "unknown":
            unknowns.append(it["text"])
            if it.get("oxy_line"):
                cur()["oxy"].append(it["oxy_line"])
            continue
        if kind == "product":
            cur()["products"].append(pair_line(it["shade"], it["amount"]))
            if it.get("oxy_line"):
                cur()["oxy"].append(it["oxy_line"])
            continue
        if kind == "oxidant":
            if it.get("text"):
                unknowns.append(it["text"])
            if it.get("oxy_line"):
                cur()["oxy"].append(it["oxy_line"])
            continue
        if kind == "shade":
            if it.get("header"):
                start_section(it["header"])
            name = it["shade"]
            if it.get("series"):
                name = f"{it['series']} {name}"
            if name:
                cur()["shades"].append(pair_line(name, it["amount"]))
            elif it["amount"]:
                cur()["shades"].append(it["amount"])
                flags.append("כותרת עם כמות בלי גוון")
            if it.get("oxy_line"):
                cur()["oxy"].append(it["oxy_line"])

    i = 1
    while i < len(sections):
        sec, prev = sections[i], sections[i - 1]
        empty = not (sec["shades"] or sec["products"] or sec["oxy"])
        prev_has = prev["shades"] or prev["products"] or prev["oxy"]
        if sec["header"] and empty and prev_has and not prev["header"]:
            prev["header"] = sec["header"]
            flags.append("כותרת שויכה אחורה לפי מיקום")
            sections.pop(i)
            continue
        i += 1

    first = True
    for sec in sections:
        if not (sec["header"] or sec["shades"] or sec["products"] or sec["oxy"]):
            continue
        if not first:
            lines.append("")
        first = False
        if sec["header"]:
            lines.append(sec["header"])
        lines.extend(drop_consecutive_dupes(sec["shades"]))
        lines.extend(drop_consecutive_dupes(sec["products"]))
        lines.extend(unique(sec["oxy"]))

    if unknowns:
        if lines and lines[-1] != "":
            lines.append("")
        for text in drop_consecutive_dupes(unknowns):
            lines.append(f"{text}  [לבדיקה]")
            flags.append("טקסט לא מזוהה")

    if instructions:
        if lines and lines[-1] != "":
            lines.append("")
        lines.append("הנחיות עבודה")
        lines.extend(instructions)

    while lines and lines[-1] == "":
        lines.pop()
    return "\n".join(lines), unique(flags), unique(corr)


def unique(seq: list[str]) -> list[str]:
    seen = set()
    out = []
    for x in seq:
        if x and x not in seen:
            seen.add(x)
            out.append(x)
    return out


def drop_consecutive_dupes(seq: list[str]) -> list[str]:
    out: list[str] = []
    for x in seq:
        if not x:
            continue
        if out and out[-1] == x:
            continue
        out.append(x)
    return out


def note_fingerprint(note: str) -> str:
    return (note or "").replace(RLM, "").replace("\u200e", "").strip()


def dedupe_identical_formulas(formulas: list[dict]) -> list[dict]:
    """Keep one note per customer when two slots are the same mix."""
    groups: dict[str, list[int]] = defaultdict(list)
    for i, f in enumerate(formulas):
        key = note_fingerprint(f.get("note") or "")
        if key:
            groups[key].append(i)
    drop: set[int] = set()
    for indexes in groups.values():
        if len(indexes) < 2:
            continue
        numbered = [i for i in indexes if formulas[i].get("f") is not None]
        keep = numbered[0] if numbered else indexes[0]
        for i in indexes:
            if i != keep:
                drop.add(i)
    return [f for i, f in enumerate(formulas) if i not in drop]


def row_signature(row: dict) -> tuple[str, str, str, str]:
    return (
        clean(row.get("shade")),
        clean(row.get("amt")),
        clean(row.get("oxy")),
        clean(row.get("oxyAmt")),
    )


def drop_duplicate_source_rows(rows: list[dict]) -> list[dict]:
    out = []
    prev = None
    for row in rows:
        sig = row_signature(row)
        if sig == prev:
            continue
        prev = sig
        out.append(row)
    return out


def load_customers() -> dict:
    wb = load_workbook(CUST_PATH, data_only=True)
    ws = wb.active
    people = {}
    for r in range(2, ws.max_row + 1):
        cid = ws.cell(r, 1).value
        if cid is None:
            continue
        people[int(cid) if isinstance(cid, (int, float)) else cid] = {
            "cust_no": int(cid) if isinstance(cid, (int, float)) else cid,
            "first": clean(ws.cell(r, 2).value),
            "last": clean(ws.cell(r, 3).value),
            "address": clean(ws.cell(r, 4).value),
            "city": clean(ws.cell(r, 5).value),
            "phone1": clean(ws.cell(r, 7).value),
            "phone2": clean(ws.cell(r, 8).value),
        }
    wb.close()
    return people


def load_color_rows() -> list[dict]:
    wb = load_workbook(COLOR_PATH, data_only=True)
    ws = wb.active
    rows = []
    for r in range(2, ws.max_row + 1):
        cust = ws.cell(r, 2).value
        if cust is None:
            continue
        rows.append({
            "src": r,
            "seq": ws.cell(r, 1).value,
            "cust": int(cust) if isinstance(cust, (int, float)) else cust,
            "f": ws.cell(r, 3).value,
            "shade": ws.cell(r, 4).value,
            "amt": ws.cell(r, 5).value,
            "oxy": ws.cell(r, 6).value,
            "oxyAmt": ws.cell(r, 7).value,
        })
    wb.close()
    return rows


def formula_sort_key(formulas: list[dict]) -> list[dict]:
    hadash_ids = {f["f"] for f in formulas if f["has_hadash"]}
    if len(hadash_ids) > 1:
        for f in formulas:
            f["flags"].append("שתי פורמולות מסומנות חדש — נבחר לפי 1→2→3 אחרי החדשות")
        newest = sorted(hadash_ids, key=lambda x: (x is None, x if isinstance(x, (int, float)) else 99))
        rest = [f for f in formulas if f["f"] not in hadash_ids]
        head = [f for n in newest for f in formulas if f["f"] == n]
        rest_sorted = sorted(rest, key=lambda f: (f["f"] is None, f["f"] if isinstance(f["f"], (int, float)) else 99))
        return head + rest_sorted
    if hadash_ids:
        newest = next(iter(hadash_ids))
        head = [f for f in formulas if f["f"] == newest]
        rest = sorted(
            [f for f in formulas if f["f"] != newest],
            key=lambda f: (f["f"] is None, f["f"] if isinstance(f["f"], (int, float)) else 99),
        )
        return head + rest
    return sorted(
        formulas,
        key=lambda f: (f["f"] is None, f["f"] if isinstance(f["f"], (int, float)) else 99),
    )


def build_customer_formulas(color_rows: list[dict]) -> dict[int, list[dict]]:
    grouped = defaultdict(lambda: defaultdict(list))
    for row in color_rows:
        f = row["f"] if row["f"] not in (None, "") else None
        grouped[row["cust"]][f].append(row)

    out = {}
    for cid, by_f in grouped.items():
        formulas = []
        for fnum, rows in by_f.items():
            rows = drop_duplicate_source_rows(rows)
            items = [classify_row(r) for r in rows]
            items = [it for it in items if it["kind"] != "empty"]
            if not items:
                continue
            attach_series(items)
            infer_missing_amounts(items)
            has_hadash = any(it["kind"] == "hadash" for it in items)
            client_notes_local = unique([it["text"] for it in items if it["kind"] == "client_note"])
            price_like = formula_is_price_like(items)
            flags = []
            if fnum is None:
                flags.append("פורמולה בלי מספר שיוך")
            note, nflags, corr = render_formula(items, [], price_like)
            raw = "\n".join(raw_line(r) for r in rows if any(clean(r[k]) for k in ("shade", "amt", "oxy", "oxyAmt")))
            formulas.append({
                "f": int(fnum) if isinstance(fnum, (int, float)) else fnum,
                "has_hadash": has_hadash,
                "is_color": formula_is_color(items),
                "price_like": price_like,
                "local_notes": client_notes_local,
                "note": note,
                "flags": unique(flags + nflags),
                "corr": unique(corr),
                "raw": raw,
                "items": items,
            })
        out[cid] = formulas
    return out


def eligibility(person: dict | None, cid, has_formulas: bool) -> tuple[str, list[str]]:
    """Return ('official'|'verify', reasons)."""
    reasons = []
    if person is None:
        return "verify", ["לקוחה לא מזוהה"]
    first, last = person["first"], person["last"]
    if not first:
        reasons.append("חסר שם פרטי")
    if not last:
        reasons.append("חסר שם משפחה")
    p1, p2 = person["phone1"], person["phone2"]
    mobile, extra = pick_mobile(p1, p2)
    reasons.extend(extra)
    if not mobile:
        k1, k2 = phone_kind(p1), phone_kind(p2)
        if not p1 and not p2:
            reasons.append("אין טלפון")
        elif "landline" in (k1, k2) and "mobile" not in (k1, k2):
            reasons.append("יש רק טלפון קווי")
        elif "mobile_uncertain" in (k1, k2):
            reasons.append("נייד לא ניתן לנרמל ל-10 ספרות תקינות")
        else:
            reasons.append("אין טלפון נייד תקין")
        return "verify", reasons
    if reasons and any(r.startswith("חסר") for r in reasons):
        return "verify", [r for r in reasons if not r.startswith("שני ניידים")]
    return "official", extra


def write_sheet(
    wb: Workbook,
    title: str,
    headers: list[str],
    rows: list[list],
    widths: list[int],
    wrap_cols: set[int],
    group_col: int | None = None,
):
    ws = wb.create_sheet(title)
    ws.sheet_view.rightToLeft = True
    ws.freeze_panes = "A2"
    ws.print_title_rows = "1:1"
    for col, header in enumerate(headers, 1):
        cell = ws.cell(1, col, header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN
    ws.row_dimensions[1].height = 28
    group_index = 0
    sentinel = object()
    prev_group = sentinel
    flag_idx = next((i for i, h in enumerate(headers) if "בדיקה" in h or "סימון" in h), None)
    for r_i, row in enumerate(rows, 2):
        group_key = row[group_col] if group_col is not None else r_i
        if group_key != prev_group:
            if prev_group is not sentinel:
                group_index += 1
            prev_group = group_key
        band = CUST_FILL_A if group_index % 2 == 0 else CUST_FILL_B
        for c_i, value in enumerate(row, 1):
            cell = ws.cell(r_i, c_i, value if value not in ("", None) else None)
            cell.font = CELL_FONT
            cell.border = THIN
            cell.fill = band
            if c_i in wrap_cols:
                cell.alignment = WRAP
            else:
                cell.alignment = RIGHT if c_i > 1 else CENTER
        ws.row_dimensions[r_i].height = fitted_row_height(row, wrap_cols, widths)
        if flag_idx is not None and row[flag_idx]:
            ws.cell(r_i, flag_idx + 1).fill = WARN_FILL
    for col, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = w
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows)+1}"
    return ws


def main():
    people = load_customers()
    color_rows = load_color_rows()
    formulas_by_cust = build_customer_formulas(color_rows)
    color_ids = set(formulas_by_cust)
    all_ids = set(people) | color_ids

    official_people = []
    verify_people = []
    phone_to_custs = defaultdict(list)
    name_phone_to_custs = defaultdict(list)

    meta = {}
    for cid in sorted(all_ids, key=lambda x: (isinstance(x, str), x)):
        person = people.get(cid)
        has_f = bool(formulas_by_cust.get(cid))
        status, reasons = eligibility(person, cid, has_f)
        mobile, extra = pick_mobile(person["phone1"], person["phone2"]) if person else (None, [])
        flags = unique(reasons if status == "verify" else extra)
        client_notes = unique(
            note
            for f in formulas_by_cust.get(cid, [])
            for note in f.get("local_notes", [])
        )
        formulas = formulas_by_cust.get(cid, [])
        for f in formulas:
            # prepend client-level notes
            if client_notes:
                body = f["note"]
                prefix = "\n".join(client_notes) + ("\n\n" if body else "")
                # avoid duplicating if render already empty of notes
                if not body.startswith(client_notes[0]):
                    f["note"] = prefix + body
        formulas = dedupe_identical_formulas(formulas)
        color_count = sum(1 for f in formulas if f["is_color"])
        if color_count >= 2 and not any(f["has_hadash"] for f in formulas):
            flags.append("כמה מיקסי צבע אמיתיים בלי סימון חדש")
        ordered = formula_sort_key(formulas) if formulas else []
        rec = {
            "cid": cid,
            "person": person,
            "status": status,
            "reasons": reasons,
            "mobile": mobile,
            "flags": flags,
            "client_notes": client_notes,
            "formulas": ordered,
        }
        meta[cid] = rec
        if status == "official":
            official_people.append(rec)
            phone_to_custs[mobile].append(cid)
            key = f"{person['first']}|{person['last']}|{mobile}"
            name_phone_to_custs[key].append(cid)
        else:
            verify_people.append(rec)

    dup_phones = {p for p, ids in phone_to_custs.items() if len(set(ids)) > 1}
    dup_people = {cid for ids in name_phone_to_custs.values() if len(set(ids)) > 1 for cid in ids}

    official_rows = []
    verify_rows = []
    review_rows = []

    def label_for(f: dict) -> str:
        if f["f"] is None:
            base = "פורמולה ללא מספר"
        else:
            base = f"פורמולה {f['f']}"
        if f["has_hadash"]:
            base += " - חדש"
        return base

    def push_review(sheet, cid, phone, formula_label, kind, detail, original, after):
        review_rows.append([
            sheet, cid, phone, formula_label, kind, detail, original, after,
        ])

    for rec in official_people:
        person, cid, mobile = rec["person"], rec["cid"], rec["mobile"]
        dup_p = "כן" if mobile in dup_phones else ""
        dup_n = "כן" if cid in dup_people else ""
        if dup_p:
            rec["flags"].append("כפילות טלפון")
        if dup_n:
            rec["flags"].append("כפילות לקוחה (שם מלא + טלפון, CustNo שונה)")
        if not rec["formulas"]:
            official_rows.append([
                mobile, person["first"], person["last"], cid, "", "אין פורמולות",
                "", "", "\n".join(rec["flags"]), "",
                person["phone1"], person["phone2"], dup_p, dup_n,
            ])
            for flag in rec["flags"]:
                push_review("רשמי", cid, mobile, "", flag, flag, "", "")
            continue
        for f in rec["formulas"]:
            all_flags = unique(rec["flags"] + f["flags"])
            official_rows.append([
                mobile, person["first"], person["last"], cid,
                f["f"] if f["f"] is not None else "",
                label_for(f), f["note"], f["raw"],
                "\n".join(all_flags), "\n".join(f["corr"]),
                person["phone1"], person["phone2"], dup_p, dup_n,
            ])
            for flag in all_flags:
                push_review("רשמי", cid, mobile, label_for(f), flag, flag, f["raw"], f["note"])
            for c in f["corr"]:
                push_review("רשמי", cid, mobile, label_for(f), "תיקון", c, f["raw"], f["note"])

    for rec in verify_people:
        person = rec["person"]
        cid = rec["cid"]
        first = person["first"] if person else ""
        last = person["last"] if person else ""
        p1 = person["phone1"] if person else ""
        p2 = person["phone2"] if person else ""
        reason = "\n".join(unique(rec["reasons"] + rec["flags"]))
        formulas = rec["formulas"] or [{
            "f": "", "has_hadash": False, "note": "", "raw": "", "flags": [], "corr": [],
            "is_color": False, "price_like": False,
        }]
        for f in formulas:
            lab = label_for(f) if f.get("note") is not None and (f.get("raw") or f.get("note") or f.get("f") != "") else ""
            if f.get("f") == "" and not f.get("raw") and not f.get("note"):
                lab = ""
            verify_rows.append([
                reason, cid, first, last, p1, p2, rec["mobile"] or "",
                f.get("f") if f.get("f") != "" else "",
                lab, f.get("note") or "", f.get("raw") or "",
                "\n".join(unique(f.get("flags") or [])),
                "\n".join(f.get("corr") or []),
            ])
            for flag in unique(rec["reasons"] + rec["flags"] + (f.get("flags") or [])):
                push_review("אימות", cid, rec["mobile"] or "", lab, flag, flag, f.get("raw") or "", f.get("note") or "")

    wb = Workbook()
    wb.remove(wb.active)
    write_sheet(
        wb, "רשמי",
        [
            "טלפון", "שם פרטי", "שם משפחה", "CustNo", "מספר פורמולה", "תווית",
            "פתק", "מקור גולמי", "סימוני בדיקה", "תיקונים",
            "Phone1 מקורי", "Phone2 מקורי", "כפילות טלפון", "כפילות לקוחה",
        ],
        official_rows,
        [14, 16, 18, 12, 12, 24, 58, 52, 32, 32, 16, 16, 12, 14],
        {7, 8, 9, 10},
        group_col=3,
    )
    write_sheet(
        wb, "אימות לקוחות",
        [
            "סיבת אימות", "CustNo", "שם פרטי", "שם משפחה", "Phone1", "Phone2",
            "נייד שזוהה", "מספר פורמולה", "תווית", "פתק", "מקור גולמי",
            "סימוני בדיקה", "תיקונים",
        ],
        verify_rows,
        [28, 12, 16, 18, 16, 16, 14, 10, 24, 58, 52, 28, 28],
        {1, 10, 11, 12, 13},
        group_col=1,
    )
    write_sheet(
        wb, "בדיקה",
        ["גיליון", "CustNo", "טלפון", "פורמולה", "סוג", "פירוט", "מקור גולמי", "אחרי עיבוד"],
        review_rows,
        [12, 12, 14, 24, 28, 36, 52, 52],
        {6, 7, 8},
        group_col=1,
    )

    n_official_cust = len({r[3] for r in official_rows})
    n_verify_cust = len({r[1] for r in verify_rows})
    summary = [
        ["לקוחות ב-CustList", len(people)],
        ["לקוחות ב-CustColor", len(color_ids)],
        ["לקוחות בצבע ולא ברשימה", len(color_ids - set(people))],
        ["שורות רשמיות (פתקים)", len(official_rows)],
        ["לקוחות בגיליון רשמי", n_official_cust],
        ["שורות אימות", len(verify_rows)],
        ["לקוחות בגיליון אימות", n_verify_cust],
        ["שורות בדיקה", len(review_rows)],
        ["טלפונים כפולים", len(dup_phones)],
        ["לקוחות כפולות (שם+טלפון)", len(dup_people)],
    ]
    write_sheet(wb, "סיכום", ["מדד", "ערך"], summary, [40, 16], set())

    wb.save(OUT_PATH)
    print(f"wrote {OUT_PATH}")
    for row in summary:
        print(f"  {row[0]}: {row[1]}")


if __name__ == "__main__":
    main()
