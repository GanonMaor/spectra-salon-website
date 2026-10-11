#!/usr/bin/env python3
"""Export professional L'Oréal barcodes found via invoice + web research to Excel."""

import json
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "data-import/reports/loreal-professional-barcodes-web.xlsx"


def ean_ok(c: str) -> bool:
    c = str(c or "")
    if len(c) != 13 or not c.isdigit():
        return False
    s = sum(int(d) if i % 2 == 0 else int(d) * 3 for i, d in enumerate(c[:12]))
    return (10 - (s % 10)) % 10 == int(c[12])


def norm_shade(s) -> str:
    s = str(s or "").strip().upper().replace(",", ".")
    s = re.sub(r"\s+", "", s)
    return s.replace("FUNDA", "").replace("#", "")


def main():
    catalog = json.loads((ROOT / "src/data/catalog-brands/l-oreal-professionnel.json").read_text())
    findings_doc = json.loads(
        (ROOT / "data-import/reports/loreal-missing-barcodes-web-findings.json").read_text()
    )
    worklist = json.loads(
        (ROOT / "data-import/reports/loreal-missing-barcodes-worklist.json").read_text()
    )

    all_catalog_bcs = set()
    inoa_by_shade = {}
    for p in catalog:
        bc = str(p.get("barcode") or "").strip()
        if bc:
            all_catalog_bcs.add(bc)
        if p.get("series") == "NEW INOA":
            inoa_by_shade[norm_shade(p.get("shade"))] = p

    invoice_rows = [
        ("INOA DEVELOPERS", "6% 20 Vol", "3474630417830", 4, "developer", "exact"),
        ("MAJIREL / OXYDANT", "20 Vol 6%", "3474637284305", 2, "developer", "web"),
        ("DIACTIVATOR", "2.7% 9 Vol", "3474630440432", 6, "developer", "exact"),
        ("BS DEVELOPERS", "OIL 9% 30 Vol", "3474636979172", 1, "developer", "web"),
        ("BLOND STUDIO BLEACH", "9 BONDER IN", "3474637105761", 1, "bleach", "web"),
        ("NEW INOA", "3", "3474637130992", 3, "color", "exact"),
        ("NEW INOA", "4", "3474637134518", 3, "color", "exact"),
        ("NEW INOA", "4.15", "3474637128142", 3, "color", "web_new"),
        ("NEW INOA", "5", "3474637134419", 3, "color", "web_new"),
        ("NEW INOA", "5.4", "3474637133313", 3, "color", "exact"),
        ("NEW INOA", "5.8", "3474637131746", 3, "color", "web_new"),
        ("NEW INOA", "5.17", "3474637131890", 3, "color", "exact"),
        ("NEW INOA", "7.11", "3474637130237", 3, "color", "exact"),
        ("NEW INOA", "7.43", "3474637133207", 3, "color", "exact"),
        ("NEW INOA", "8", "3474637134464", 3, "color", "exact"),
        ("NEW INOA", "8.1", "3474637130183", 3, "color", "exact"),
        ("MAJIREL FUNDA", "1", "3474637283261", 3, "color", "exact"),
        ("MAJIREL FUNDA", "5", "3474637283230", 9, "color", "exact"),
        ("MAJIREL FUNDA", "5.8", "3474637280635", 3, "color", "exact"),
        ("MAJIREL FUNDA", "5.12", "3474637282721", 3, "color", "exact"),
        ("MAJIREL FUNDA", "6", "3474637283223", 9, "color", "invoice_new"),
        ("MAJIREL FUNDA", "7", "3474637283216", 9, "color", "exact"),
        ("MAJIREL FUNDA", "7.18", "3474637280994", 3, "color", "invoice_new"),
        ("MAJIREL FUNDA", "8", "3474637283209", 3, "color", "exact"),
        ("MAJIREL FUNDA", "8.3", "3474637281878", 3, "color", "exact"),
        ("MAJIREL FUNDA", "9", "3474637283193", 3, "color", "exact"),
        ("MAJIREL FUNDA", "6.1 CC", "3474637279899", 6, "color", "fixed"),
        ("MAJIREL FUNDA", "7.1 CC", "3474637279882", 6, "color", "fixed"),
        ("DIA LIGHT 60g", "Clear", "3474637320423", 6, "color", "fixed"),
        ("DIA LIGHT 60g", "6.11", "3474637322328", 6, "color", "fixed"),
        ("DIA LIGHT 60g", "6.28", "3474637322588", 3, "color", "fixed"),
        ("DIA LIGHT 60g", "7.18", "3474637322472", 3, "color", "fixed"),
        ("DIA LIGHT 60g", "7.23", "3474637321345", 3, "color", "fixed"),
        ("DIA LIGHT 60g", "7.43", "3474637321581", 3, "color", "fixed"),
        ("DIA LIGHT 60g", "8.23", "3474637321338", 3, "color", "fixed"),
        ("DIA LIGHT 60g", "9.02", "3474637320133", 6, "color", "exact"),
        ("DIA LIGHT 60g", "10.12", "3474637320669", 3, "color", "fixed"),
        ("DIA LIGHT 60g", "10.23", "3474637321321", 3, "color", "exact"),
        ("DIA LIGHT 60g", "10.32", "3474637321123", 3, "color", "exact"),
        ("DIA LIGHT", "10.82", "3474637156534", 3, "color", "invoice_new"),
        ("DIA COLOR", "3", "3474637206864", 6, "color", "exact"),
        ("DIA COLOR", "4", "3474637206871", 3, "color", "exact"),
        ("DIA COLOR", "6", "3474637206895", 3, "color", "exact"),
        ("DIA COLOR", "8.34", "3474637208110", 6, "color", "exact"),
    ]

    glama_inoa = [
        ("7.13", "3474637127916"),
        ("6.23", "3474637127961"),
        ("7.23", "3474637128029"),
        ("8.13", "3474637128074"),
        ("7.18", "3474637128234"),
        ("9.13", "3474637128289"),
        ("5.18", "3474637128401"),
        ("6.32", "3474637128555"),
        ("6.66", "3474637128838"),
        ("8.34", "3474637129439"),
        ("6.35", "3474637129804"),
        ("7.1", "3474637130114"),
        ("7.11", "3474637130213"),
        ("9.1", "3474637130268"),
        ("5.1", "3474637130312"),
        ("10.1", "3474637130367"),
        ("5.12", "3474637130411"),
        ("9.12", "3474637130473"),
        ("8.11", "3474637130527"),
        ("9.11", "3474637130572"),
        ("6.11", "3474637130626"),
        ("7.0", "3474637130879"),
        ("6.0", "3474637130923"),
        ("3", "3474637130978"),
        ("5.0", "3474637131029"),
        ("7.3 F", "3474637131074"),
        ("10", "3474637131128"),
        ("6.3 F", "3474637131173"),
        ("8.3 F", "3474637131227"),
        ("1", "3474637131272"),
        ("8.0", "3474637131326"),
        ("5.3 F", "3474637131371"),
        ("4.0", "3474637131425"),
        ("9.3 F", "3474637131470"),
        ("9.0", "3474637131524"),
        ("6.8", "3474637131678"),
        ("5.8", "3474637131722"),
        ("7.8", "3474637131777"),
        ("5.17", "3474637131876"),
        ("7.43", "3474637133184"),
        ("7.44", "3474637133238"),
        ("6.45", "3474637133351"),
        ("7.4", "3474637133412"),
        ("6.46", "3474637133474"),
        ("6", "3474637134297"),
        ("7", "3474637134341"),
        ("5", "3474637134396"),
        ("8", "3474637134440"),
        ("4", "3474637134495"),
        ("9", "3474637134549"),
    ]

    inoa_web_new = [
        ("4.15", "3474637128142", "invoice_il+EU shops", "high", "אריזה חדשה מול קטלוג 3474637122270"),
        ("5.8", "3474637131746", "invoice_il+EU shops", "high", "אריזה חדשה מול קטלוג 3474637131722"),
        ("5", "3474637134419", "invoice_il", "high", "אריזה חדשה מול קטלוג 3474637134396"),
        ("5.17", "3474637131890", "invoice_il+catalog", "high", "בישראל; Partner Shop גם 3474637131876"),
        ("6.20", "3474636727407", "probeauty+ellisons", "high", "חסר כ־6.20; קיים כ־6.2"),
        ("5.23", "", "web search", "low", "לא נמצא EAN ציבורי"),
    ]

    gap_fills = findings_doc["findings"]
    extra = [
        {
            "series": "MAJIREL FUNDA",
            "shade": "7.18",
            "barcode": "3474637280796",
            "type": "color",
            "source": "harotec_de",
            "confidence": "medium",
            "category": "regional_variant",
            "notes": "וריאנט EU ל־7.18",
            "ean_valid": True,
            "priority": "P1",
        },
        {
            "series": "MAJIREL FUNDA",
            "shade": "7.18",
            "barcode": "3474637280895",
            "type": "color",
            "source": "es_shop",
            "confidence": "medium",
            "category": "regional_variant",
            "notes": "וריאנט ES ל־7.18",
            "ean_valid": True,
            "priority": "P1",
        },
        {
            "series": "NEW INOA",
            "shade": "5.17",
            "barcode": "3474637131876",
            "type": "color",
            "source": "loreal_partner_shop",
            "confidence": "high",
            "category": "regional_variant",
            "notes": "Partner Shop AU/SG; ישראל 7131890",
            "ean_valid": True,
            "priority": "P0",
        },
    ]

    header_fill = PatternFill("solid", fgColor="1F3A5F")
    header_font = Font(bold=True, color="FFFFFF", name="Arial", size=11)
    new_fill = PatternFill("solid", fgColor="FFF3CD")
    diff_fill = PatternFill("solid", fgColor="F8D7DA")
    ok_fill = PatternFill("solid", fgColor="D4EDDA")
    thin = Border(
        left=Side(style="thin", color="CCCCCC"),
        right=Side(style="thin", color="CCCCCC"),
        top=Side(style="thin", color="CCCCCC"),
        bottom=Side(style="thin", color="CCCCCC"),
    )

    def style_header(ws, ncols):
        for col in range(1, ncols + 1):
            cell = ws.cell(1, col)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[1].height = 28
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    def autosize(ws, min_w=10, max_w=48):
        for col in ws.columns:
            letter = get_column_letter(col[0].column)
            length = 0
            for cell in col:
                if cell.value:
                    length = max(length, len(str(cell.value)))
            ws.column_dimensions[letter].width = max(min_w, min(max_w, length + 2))

    def write_rows(ws, headers, rows, highlight_fn=None):
        ws.append(headers)
        style_header(ws, len(headers))
        for r in rows:
            ws.append(list(r))
            row_idx = ws.max_row
            for c in range(1, len(headers) + 1):
                cell = ws.cell(row_idx, c)
                cell.border = thin
                cell.alignment = Alignment(vertical="center", wrap_text=True)
            if highlight_fn:
                fill = highlight_fn(r)
                if fill:
                    for c in range(1, len(headers) + 1):
                        ws.cell(row_idx, c).fill = fill
        autosize(ws)

    def find_inoa_cat(shade: str):
        ns = norm_shade(shade)
        cat = inoa_by_shade.get(ns)
        if cat:
            return cat
        if ns == "6.20":
            return inoa_by_shade.get("6.2")
        if ns.endswith("F"):
            return inoa_by_shade.get(ns[:-1]) or inoa_by_shade.get(ns.replace("F", ""))
        for k, v in inoa_by_shade.items():
            if k.replace("F", "") == ns.replace("F", ""):
                return v
        return None

    wb = Workbook()

    # Sheet 1 — all professional
    ws1 = wb.active
    ws1.title = "כל_הברקודים_המקצועיים"
    all_rows = []
    seen_bc = set()
    for f in gap_fills + extra:
        bc = f.get("barcode") or ""
        if not bc or bc in seen_bc:
            continue
        seen_bc.add(bc)
        in_cat = "כן" if bc in all_catalog_bcs else "לא — חדש לנו"
        all_rows.append(
            [
                f.get("priority", ""),
                f.get("category", ""),
                f.get("series", ""),
                f.get("shade", ""),
                bc,
                f.get("type", ""),
                f.get("confidence", ""),
                "כן" if f.get("ean_valid") or ean_ok(bc) else "לא",
                in_cat,
                f.get("source", ""),
                f.get("notes", ""),
            ]
        )
    for ser, sh, bc, qty, typ, st in invoice_rows:
        if bc in seen_bc:
            continue
        seen_bc.add(bc)
        in_cat = "כן" if bc in all_catalog_bcs else "לא — חדש לנו"
        all_rows.append(
            [
                "P0" if in_cat.startswith("לא") else "P3",
                "invoice_order",
                ser,
                sh,
                bc,
                typ,
                st,
                "כן" if ean_ok(bc) else "לא",
                in_cat,
                "תעודת לוריאל ישראל",
                f"כמות בהזמנה: {qty}",
            ]
        )
    # Glama diffs as professional findings too
    for shade, bc in glama_inoa:
        if bc in seen_bc:
            continue
        cat = find_inoa_cat(shade)
        cat_bc = str((cat or {}).get("barcode") or "")
        if cat_bc == bc:
            continue
        seen_bc.add(bc)
        in_cat = "כן" if bc in all_catalog_bcs else "לא — חדש לנו"
        all_rows.append(
            [
                "P1",
                "inoa_regional_variant",
                "NEW INOA",
                shade,
                bc,
                "color",
                "medium",
                "כן" if ean_ok(bc) else "לא",
                in_cat,
                "glama.co",
                f"וריאנט מול קטלוג {cat_bc}" if cat_bc else "לא בקטלוג",
            ]
        )

    all_rows.sort(key=lambda r: (0 if str(r[8]).startswith("לא") else 1, str(r[2]), str(r[3])))
    write_rows(
        ws1,
        [
            "עדיפות",
            "קטגוריה",
            "סדרה",
            "גוון",
            "ברקוד",
            "סוג",
            "ביטחון",
            "EAN תקין",
            "בקטלוג שלנו?",
            "מקור",
            "הערות",
        ],
        all_rows,
        lambda r: new_fill if str(r[8]).startswith("לא") else None,
    )

    # Sheet 2 — INOA focus
    ws2 = wb.create_sheet("INOA_ברקודים_חדשים")
    inoa_rows = []
    for shade, bc, source, conf, notes in inoa_web_new:
        cat = find_inoa_cat(shade)
        cat_bc = str((cat or {}).get("barcode") or "")
        if not bc:
            status = "לא נמצא ברשת"
            eanv = ""
        elif not cat_bc:
            status = "חסר בקטלוג — להוסיף"
            eanv = "כן" if ean_ok(bc) else "לא"
        elif cat_bc == bc:
            status = "כבר בקטלוג"
            eanv = "כן" if ean_ok(bc) else "לא"
        else:
            status = "ברקוד חדש (שונה מהקטלוג)"
            eanv = "כן" if ean_ok(bc) else "לא"
        inoa_rows.append(
            ["WEB/INVOICE", "NEW INOA", shade, bc, cat_bc, status, conf, source, notes, eanv]
        )

    for shade, bc in glama_inoa:
        cat = find_inoa_cat(shade)
        cat_bc = str((cat or {}).get("barcode") or "")
        if cat_bc == bc:
            continue
        status = "חסר בקטלוג" if not cat_bc else "וריאנט אזורי / ברקוד שונה"
        inoa_rows.append(
            [
                "GLAMA_AU",
                "NEW INOA",
                shade,
                bc,
                cat_bc,
                status,
                "medium",
                "glama.co",
                "ייתכן אריזה/שוק שונה — לא לדרוס בלי אימות",
                "כן" if ean_ok(bc) else "לא",
            ]
        )

    for ser, sh, bc, qty, typ, st in invoice_rows:
        if "INOA" not in ser:
            continue
        if any(r[3] == bc for r in inoa_rows):
            continue
        cat = find_inoa_cat(sh)
        cat_bc = str((cat or {}).get("barcode") or "")
        status = "כבר בקטלוג" if bc in all_catalog_bcs else "חדש מהתעודה"
        inoa_rows.append(
            [
                "INVOICE_IL",
                "NEW INOA" if ser == "NEW INOA" else ser,
                sh,
                bc,
                cat_bc,
                status,
                "high",
                "תעודת ישראל",
                f"qty={qty}",
                "כן" if ean_ok(bc) else "לא",
            ]
        )

    rank = {
        "ברקוד חדש (שונה מהקטלוג)": 0,
        "חסר בקטלוג — להוסיף": 0,
        "חסר בקטלוג": 0,
        "חדש מהתעודה": 0,
        "וריאנט אזורי / ברקוד שונה": 1,
        "לא נמצא ברשת": 2,
        "כבר בקטלוג": 3,
    }
    inoa_rows.sort(key=lambda r: (rank.get(r[5], 9), str(r[2])))

    def hl_inoa(r):
        if r[5] in (
            "ברקוד חדש (שונה מהקטלוג)",
            "חסר בקטלוג — להוסיף",
            "חסר בקטלוג",
            "חדש מהתעודה",
        ):
            return new_fill
        if r[5] == "וריאנט אזורי / ברקוד שונה":
            return diff_fill
        if r[5] == "כבר בקטלוג":
            return ok_fill
        return None

    write_rows(
        ws2,
        [
            "מקור",
            "סדרה",
            "גוון",
            "ברקוד_שנמצא",
            "ברקוד_בקטלוג",
            "סטטוס",
            "ביטחון",
            "מקור_פירוט",
            "הערות",
            "EAN_תקין",
        ],
        inoa_rows,
        hl_inoa,
    )

    # Sheet 3 — invoice
    ws3 = wb.create_sheet("תעודת_ישראל")
    inv = [
        [
            ser,
            sh,
            bc,
            qty,
            typ,
            st,
            "כן" if bc in all_catalog_bcs else "לא — חדש",
            "כן" if ean_ok(bc) else "לא",
        ]
        for ser, sh, bc, qty, typ, st in invoice_rows
    ]
    write_rows(
        ws3,
        ["סדרה", "גוון", "ברקוד", "כמות", "סוג", "סטטוס_אימות", "בקטלוג?", "EAN_תקין"],
        inv,
        lambda r: new_fill if str(r[6]).startswith("לא") else ok_fill,
    )

    # Sheet 4 — Majirel Funda
    ws4 = wb.create_sheet("Majirel_Funda")
    funda = []
    for f in gap_fills + extra:
        ser = f.get("series", "")
        if "MAJIREL" in ser.upper() or "FUNDA" in ser.upper() or "COOL" in ser.upper():
            funda.append(
                [
                    ser,
                    f.get("shade"),
                    f.get("barcode"),
                    f.get("confidence"),
                    "כן" if f.get("barcode") in all_catalog_bcs else "לא — חדש",
                    f.get("source"),
                    f.get("notes"),
                ]
            )
    for ser, sh, bc, qty, typ, st in invoice_rows:
        if "MAJIREL" not in ser:
            continue
        if any(x[2] == bc for x in funda):
            continue
        funda.append(
            [
                ser,
                sh,
                bc,
                st,
                "כן" if bc in all_catalog_bcs else "לא — חדש",
                "תעודת ישראל",
                f"qty={qty}",
            ]
        )
    funda.sort(key=lambda r: (0 if str(r[4]).startswith("לא") else 1, str(r[1])))
    write_rows(
        ws4,
        ["סדרה", "גוון", "ברקוד", "ביטחון", "בקטלוג?", "מקור", "הערות"],
        funda,
        lambda r: new_fill if str(r[4]).startswith("לא") else ok_fill,
    )

    # Sheet 5 — still missing
    ws5 = wb.create_sheet("עדיין_חסר")
    miss = []
    for x in findings_doc["still_missing"]:
        miss.append([x["series"], x["shade"], x["type"], x["notes"]])
    found_keys = set()
    for f in gap_fills + extra:
        if f.get("barcode"):
            found_keys.add((f.get("series"), str(f.get("shade"))))
    found_keys |= {
        ("NEW INOA", "6.20"),
        ("MAJIREL FUNDA", "7.18"),
        ("DIA LIGHT 60g", "9.82"),
        ("COVER 5", "6"),
        ("COVER 5", "7"),
        ("DIA RICHESSE", "6.40"),
        ("MAJIREL", "6.62"),
        ("INOA GLOW", "L,18"),
        ("INOA GLOW", "D.18 / .18"),
    }
    for w in worklist:
        key = (w["series"], str(w["shade"]))
        found = key in found_keys
        if w["series"] == "NEW INOA" and w["shade"] == "6.20":
            found = True
        if w["series"] == "MAJIREL FUNDA" and w["shade"] in {
            "4.15",
            "4.20",
            "4.3",
            "4.35",
            "4.56",
            "6.8",
            "7.18",
            "7.4",
            "10.21",
        }:
            found = True
        if w["series"] == "DIA LIGHT 60g" and w["shade"] == "9.82":
            found = True
        if w["series"] == "COVER 5" and w["shade"] in ("6", "7"):
            found = True
        if w["series"] == "DIA RICHESSE" and w["shade"] == "6.40":
            found = True
        if w["series"] == "INOA GLOW" and "18" in str(w["shade"]):
            found = True
        if w["series"] == "MAJIREL" and w["shade"] == "6.62":
            found = True
        if not found and w["type"] == "color":
            miss.append(
                [w["series"], w["shade"], w["type"], "עדיין בלי ברקוד בקטלוג — לא נמצא בחיפוש זה"]
            )
    seen = set()
    miss2 = []
    for m in miss:
        k = (m[0], m[1])
        if k in seen:
            continue
        seen.add(k)
        miss2.append(m)
    write_rows(ws5, ["סדרה", "גוון", "סוג", "הערות"], miss2)

    # Sheet 6 — legend
    ws6 = wb.create_sheet("מקרא")
    ws6["A1"] = "מקרא — ברקודים מקצועיים לוריאל"
    ws6["A1"].font = Font(bold=True, size=14, name="Arial")
    ws6.append(["צבע/קוד", "משמעות"])
    style_header(ws6, 2)
    legend = [
        ("צהוב", "ברקוד חדש לנו — לא קיים בקטלוג; להוסיף אחרי אימות"),
        ("אדום בהיר", "וריאנט אזורי (למשל AU) ששונה מהברקוד אצלנו — לא לדרוס אוטומטית"),
        ("ירוק", "כבר קיים בקטלוג"),
        ("P0", "עדיפות גבוהה — מתעודה / חסר בקטלוג"),
        ("P1", "השלמת פער / וריאנט לצבעי מספרה"),
        ("347463712/713", "New INOA"),
        ("347463728", "New Majirel / Majirel Funda 60ml"),
        ("347463732", "Dia Light 60g"),
        ("INOA חשוב", "מהתעודה: 4.15=3474637128142 · 5.8=3474637131746 · 5=3474637134419"),
    ]
    for a, b in legend:
        ws6.append([a, b])
    ws6["A3"].fill = new_fill
    ws6["A4"].fill = diff_fill
    ws6["A5"].fill = ok_fill
    autosize(ws6, max_w=90)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"Wrote {OUT}")
    print(f"all_rows={len(all_rows)} new={sum(1 for r in all_rows if str(r[8]).startswith('לא'))}")
    print(
        f"inoa_rows={len(inoa_rows)} new_or_diff={sum(1 for r in inoa_rows if r[5] not in ('כבר בקטלוג',))}"
    )


if __name__ == "__main__":
    main()
