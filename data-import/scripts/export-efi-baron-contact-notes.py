#!/usr/bin/env python3
"""One contact row per official customer, with all formulas in a single note.

Headers inside the note: פורמולה חדשה / פורמולה ישנה / מחיר / שירות.
Same people as EfiBaron-contacts.xlsx (name + valid mobile; Phone2 if both mobiles).
"""

from __future__ import annotations

import importlib.util
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

SRC_DIR = Path("/Users/maorganon/Downloads/EfiBaronOld-sys")
CUST_PATH = SRC_DIR / "CustList.xlsx"
OUT_PATH = SRC_DIR / "EfiBaron-contact-notes.xlsx"

HELPERS = Path(__file__).with_name("build-efi-baron-formula-notes.py")
spec = importlib.util.spec_from_file_location("efi_notes", HELPERS)
efi = importlib.util.module_from_spec(spec)
spec.loader.exec_module(efi)

HEADER_FILL = PatternFill("solid", fgColor="1F4E5F")
HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=11)
CELL_FONT = Font(name="Arial", size=11)
BAND_A = PatternFill("solid", fgColor="FFFFFF")
BAND_B = PatternFill("solid", fgColor="E8EEF4")
THIN = Border(
    left=Side(style="thin", color="D0D7DE"),
    right=Side(style="thin", color="D0D7DE"),
    top=Side(style="thin", color="D0D7DE"),
    bottom=Side(style="thin", color="D0D7DE"),
)
RIGHT = Alignment(horizontal="right", vertical="center", readingOrder=2)
LEFT = Alignment(horizontal="left", vertical="center")
WRAP = Alignment(horizontal="right", vertical="top", wrap_text=True, readingOrder=2)

HEADERS = ["שם פרטי", "שם משפחה", "מספר טלפון", "אימייל", "פתק"]
WIDTHS = [18, 22, 16, 32, 58]


def pick_mobile(phone1: str, phone2: str) -> str | None:
    k1, k2 = efi.phone_kind(phone1), efi.phone_kind(phone2)
    n1, n2 = efi.national_phone(phone1), efi.national_phone(phone2)
    m1 = n1 if k1 == "mobile" else None
    m2 = n2 if k2 == "mobile" else None
    if m1 and m2:
        return m2
    return m2 or m1


def clean_email(raw) -> str:
    s = efi.clean(raw)
    if not s or "@" not in s:
        return ""
    return s.strip()


def color_header(index: int, total: int) -> str:
    if total == 1:
        return "פורמולה"
    if index == 0:
        return "פורמולה חדשה"
    if index == 1:
        return "פורמולה ישנה"
    return f"פורמולה ישנה {index}"


def other_header(formula: dict) -> str:
    if formula.get("price_like"):
        return "מחיר / שירות"
    return "הערה"


def combine_note(formulas: list[dict], client_notes: list[str]) -> str:
    ordered = efi.formula_sort_key(efi.dedupe_identical_formulas(formulas))
    color = [f for f in ordered if f.get("is_color")]
    other = [f for f in ordered if not f.get("is_color")]

    parts: list[str] = []
    if client_notes:
        parts.extend(client_notes)
        parts.append("")

    def add_block(header: str, formula: dict) -> None:
        body = (formula.get("note") or "").strip()
        if not body:
            return
        if parts and parts[-1] != "":
            parts.append("")
        parts.append(header)
        parts.append(body)

    for i, formula in enumerate(color):
        add_block(color_header(i, len(color)), formula)
    for formula in other:
        add_block(other_header(formula), formula)

    while parts and parts[-1] == "":
        parts.pop()
    return "\n".join(parts)


def load_people() -> list[dict]:
    wb = load_workbook(CUST_PATH, data_only=True)
    ws = wb.active
    people = []
    for r in range(2, ws.max_row + 1):
        cid = ws.cell(r, 1).value
        if cid is None:
            continue
        first = efi.clean(ws.cell(r, 2).value)
        last = efi.clean(ws.cell(r, 3).value)
        if not first or not last:
            continue
        mobile = pick_mobile(efi.clean(ws.cell(r, 7).value), efi.clean(ws.cell(r, 8).value))
        if not mobile:
            continue
        people.append({
            "cid": int(cid) if isinstance(cid, (int, float)) else cid,
            "first": first,
            "last": last,
            "mobile": mobile,
            "email": clean_email(ws.cell(r, 11).value),
        })
    wb.close()
    people.sort(key=lambda p: (p["last"], p["first"], p["mobile"]))
    return people


def main() -> None:
    people = load_people()
    formulas_by_cust = efi.build_customer_formulas(efi.load_color_rows())

    rows = []
    with_note = 0
    with_email = 0
    for person in people:
        formulas = formulas_by_cust.get(person["cid"], [])
        client_notes = efi.unique(
            note for f in formulas for note in f.get("local_notes", [])
        )
        note = combine_note(formulas, client_notes)
        if note:
            with_note += 1
        if person["email"]:
            with_email += 1
        rows.append([
            person["first"],
            person["last"],
            person["mobile"],
            person["email"],
            note,
        ])

    wb = Workbook()
    ws = wb.active
    ws.title = "לקוחות"
    ws.sheet_view.rightToLeft = True
    ws.freeze_panes = "A2"
    for col, header in enumerate(HEADERS, 1):
        cell = ws.cell(1, col, header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN
        ws.column_dimensions[get_column_letter(col)].width = WIDTHS[col - 1]
    ws.row_dimensions[1].height = 26

    wrap_cols = {5}
    for i, row in enumerate(rows, 2):
        fill = BAND_A if i % 2 == 0 else BAND_B
        for c, value in enumerate(row, 1):
            cell = ws.cell(i, c, value or None)
            cell.font = CELL_FONT
            cell.border = THIN
            cell.fill = fill
            if c == 5:
                cell.alignment = WRAP
            elif c == 4:
                cell.alignment = LEFT
            else:
                cell.alignment = RIGHT
        ws.row_dimensions[i].height = efi.fitted_row_height(row, wrap_cols, WIDTHS)

    ws.auto_filter.ref = f"A1:E{max(1, len(rows) + 1)}"
    wb.save(OUT_PATH)
    print(f"wrote {OUT_PATH}")
    print(f"  לקוחות: {len(rows)}")
    print(f"  עם פתק פורמולה: {with_note}")
    print(f"  בלי פורמולה: {len(rows) - with_note}")
    print(f"  עם אימייל: {with_email}")


if __name__ == "__main__":
    main()
